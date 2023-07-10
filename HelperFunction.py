import PyPDF2, openpyxl, xlrd, zipfile, re
from pptx import Presentation

def docxTotalPage(path):
    try:
        with zipfile.ZipFile(path, 'r') as docx_zip:
            xml_content = docx_zip.read('docProps/app.xml')
            ptn_rslt = re.search('<Pages>(.*)</Pages>', xml_content.decode())
            if ptn_rslt:
                return int(ptn_rslt.group(1))
    except (KeyError, zipfile.BadZipFile):
        pass
    return 0

def pdfTotalPage(path):
    try:
        with open(path, "rb") as file:
            pdf_reader = PyPDF2.PdfReader(file)
            num_pages = len(pdf_reader.pages)
            return num_pages
    except PyPDF2.errors.FileNotDecryptedError:
        pass
    return 0

def pptTotalPage(path):
    presentation = Presentation(path)
    return len(presentation.slides)

def xlxTotalSheet(path):
    try:  
        workbook = openpyxl.load_workbook(path)
        num_sheets = len(workbook.sheetnames)
    except:
        workbook = xlrd.open_workbook(path)
        num_sheets = workbook.nsheets
    return num_sheets